"""Import NOAA California strandings as coarse counts, with no row-level storage.

Run locally against a database you control::

    PYTHONPATH=backend python -m lifejacket.services.historical_strandings \
        /private/Zhang_CA_2006-2025.xlsx

The workbook and its detailed locality, identifiers, and remarks never enter
the application database. Re-import atomically replaces this dataset's cells.
"""

from __future__ import annotations

import argparse
import hashlib
import math
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import delete
from sqlalchemy.orm import Session

from lifejacket.models.tables import HistoricalStrandingCellRow, HistoricalStrandingDatasetRow

DATASET_ID = "noaa_ca_2006_2025"
GRID_DEGREES = 0.1  # approximately 11 km north-south; not a point location
MIN_PUBLIC_CELL_COUNT = 10
REQUIRED_COLUMNS = (
    "State", "Latitude", "Longitude", "Latitude Units", "Longitude Units",
    "Observation Date", "Confidence Code",
)


def _decimal_coordinate(value: object, units: object) -> float | None:
    # DMS fields are mixed with decimal fields in the source; do not silently
    # interpret a DMS string as a decimal coordinate.
    if str(units).strip().lower() not in {"decimal degrees", "dec deg"}:
        return None
    try:
        coordinate = float(value)
    except (TypeError, ValueError):
        return None
    return coordinate if math.isfinite(coordinate) else None


def aggregate_workbook(path: Path) -> tuple[Counter[tuple[int, int]], Counter[str]]:
    """Stream the source; return counts and mutually exclusive exclusion reasons."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError("Expected exactly one NOAA data sheet")
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        columns = {str(name).strip(): i for i, name in enumerate(header) if name}
        missing = set(REQUIRED_COLUMNS) - columns.keys()
        if missing:
            raise ValueError(f"Missing NOAA columns: {', '.join(sorted(missing))}")

        cells: Counter[tuple[int, int]] = Counter()
        stats: Counter[str] = Counter()
        for row in rows:
            stats["total"] += 1
            get = lambda name, values=row: values[columns[name]]  # noqa: E731
            if str(get("State")).strip().lower() not in {"ca", "california"}:
                stats["outside_california_state"] += 1
                continue
            observed = get("Observation Date")
            if not isinstance(observed, (date, datetime)) or not 2006 <= observed.year <= 2025:
                stats["invalid_date"] += 1
                continue
            confidence = str(get("Confidence Code")).strip().lower()
            if "unconfirmed" in confidence:
                stats["unconfirmed"] += 1
                continue
            if not confidence.startswith("confirmed"):
                stats["unknown_confidence"] += 1
                continue
            lat = _decimal_coordinate(get("Latitude"), get("Latitude Units"))
            lon = _decimal_coordinate(get("Longitude"), get("Longitude Units"))
            if lat is None or lon is None:
                stats["invalid_or_nondecimal_coordinates"] += 1
                continue
            # Broad bounding box only. The State field is checked separately;
            # boundary and coastline validation requires a spatial dataset.
            if not (32 <= lat <= 42.5 and -125 <= lon <= -114):
                stats["outside_california_bounds"] += 1
                continue
            cells[(math.floor(lat / GRID_DEGREES), math.floor(lon / GRID_DEGREES))] += 1
            stats["included"] += 1
        return cells, stats
    finally:
        workbook.close()


def import_workbook(path: Path, session: Session) -> dict[str, object]:
    """Validate before changing the DB; replace existing aggregate in one transaction."""
    path = path.expanduser().resolve(strict=True)
    with path.open("rb") as source:
        digest = hashlib.file_digest(source, "sha256").hexdigest()
    cells, stats = aggregate_workbook(path)
    if not stats["included"]:
        raise ValueError("No usable California records; existing import was not changed")

    session.execute(delete(HistoricalStrandingCellRow).where(
        HistoricalStrandingCellRow.dataset_id == DATASET_ID
    ))
    session.execute(delete(HistoricalStrandingDatasetRow).where(
        HistoricalStrandingDatasetRow.dataset_id == DATASET_ID
    ))
    session.add(HistoricalStrandingDatasetRow(
        dataset_id=DATASET_ID,
        sha256=digest,
        total_records=stats["total"],
        included_records=stats["included"],
        exclusions_json={
            key: value for key, value in stats.items() if key not in {"total", "included"}
        },
    ))
    session.flush()
    session.add_all(
        HistoricalStrandingCellRow(
            dataset_id=DATASET_ID, latitude_index=lat, longitude_index=lon, count=count
        )
        for (lat, lon), count in cells.items()
    )
    return {"sha256": digest, "total": stats["total"], "included": stats["included"],
            "excluded": {
                key: value for key, value in stats.items() if key not in {"total", "included"}
            },
            "cells": len(cells)}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    from lifejacket.models.db import init_db, session_scope

    init_db()
    with session_scope() as db:
        result = import_workbook(args.workbook, db)
    print(result)
